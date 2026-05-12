#!/usr/bin/env python3
"""
Importa usuários legados (usuarios.xlsx) e gera vínculos demanda-usuário por email.

Fontes:
- usuarios.xlsx -> coluna D (E-mail)
- legacy_pa_itens.fiscal_email (originado de Itens_Solicitados coluna N)

Uso:
  python scripts/import_legacy_users_xlsx.py
  python scripts/import_legacy_users_xlsx.py --apply
"""

from __future__ import annotations

import argparse
import json
import os
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional
from urllib.parse import urlencode

import openpyxl
import requests


def load_env_file(env_path: Path) -> Dict[str, str]:
    env: Dict[str, str] = {}
    if not env_path.exists():
        return env

    for line in env_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        env[key.strip()] = value.strip().strip('"').strip("'")
    return env


def unique_headers(raw_headers: Iterable[Any]) -> List[str]:
    seen: Dict[str, int] = {}
    result: List[str] = []
    for idx, header in enumerate(raw_headers, start=1):
        base = str(header).strip() if header is not None and str(header).strip() else f"__col_{idx}"
        if base not in seen:
            seen[base] = 1
            result.append(base)
        else:
            seen[base] += 1
            result.append(f"{base}__{seen[base]}")
    return result


def normalize_email(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip().lower()
    if not text or "@" not in text:
        return None
    return text


def to_date(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def chunked(items: List[Dict[str, Any]], size: int) -> Iterable[List[Dict[str, Any]]]:
    for i in range(0, len(items), size):
        yield items[i : i + size]


def upsert_rows(
    base_url: str,
    service_key: str,
    table: str,
    rows: List[Dict[str, Any]],
    on_conflict: str,
) -> None:
    if not rows:
        return

    endpoint = f"{base_url}/rest/v1/{table}"
    query = urlencode({"on_conflict": on_conflict})
    url = f"{endpoint}?{query}"
    headers = {
        "apikey": service_key,
        "Authorization": f"Bearer {service_key}",
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates,return=minimal",
    }

    for batch in chunked(rows, 500):
        sanitized_batch = json.loads(json.dumps(batch, ensure_ascii=False, default=str))
        resp = requests.post(url, headers=headers, json=sanitized_batch, timeout=120)
        if resp.status_code >= 300:
            raise RuntimeError(
                f"Upsert failed on table {table} ({resp.status_code}): {resp.text}"
            )


def fetch_all_legacy_items(base_url: str, service_key: str) -> List[Dict[str, Any]]:
    headers = {
        "apikey": service_key,
        "Authorization": f"Bearer {service_key}",
    }
    page_size = 1000
    offset = 0
    output: List[Dict[str, Any]] = []
    while True:
        query = (
            f"{base_url}/rest/v1/legacy_pa_itens"
            f"?select=legacy_year,demand_code,fiscal_email"
            f"&limit={page_size}&offset={offset}"
        )
        resp = requests.get(query, headers=headers, timeout=120)
        if resp.status_code >= 300:
            raise RuntimeError(
                f"Fetch failed on legacy_pa_itens ({resp.status_code}): {resp.text}"
            )
        data = resp.json()
        if not data:
            break
        output.extend(data)
        if len(data) < page_size:
            break
        offset += page_size
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--file",
        default=str((Path(__file__).resolve().parents[2] / "usuarios.xlsx")),
    )
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    env_local = load_env_file(Path(__file__).resolve().parents[1] / ".env.local")
    supabase_url = os.getenv("NEXT_PUBLIC_SUPABASE_URL") or env_local.get("NEXT_PUBLIC_SUPABASE_URL")
    service_key = os.getenv("SUPABASE_SERVICE_ROLE_KEY") or env_local.get("SUPABASE_SERVICE_ROLE_KEY")
    if not supabase_url or not service_key:
        raise RuntimeError("NEXT_PUBLIC_SUPABASE_URL e SUPABASE_SERVICE_ROLE_KEY são obrigatórios.")

    xlsx_path = Path(args.file).resolve()
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    headers = unique_headers(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        data = {headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))}
        rows.append(data)

    directory_rows_map: Dict[str, Dict[str, Any]] = {}
    for row in rows:
        email = normalize_email(row.get("E-mail"))
        if not email:
            continue

        server_name = (
            str(row.get("Servidor") or "").strip()
            or str(row.get("Servidor__2") or "").strip()
            or None
        )
        payload = {
            "email": email,
            "server_name": server_name,
            "birth_date": to_date(row.get("Data de aniversário")),
            "start_date": to_date(row.get("Data de inicio")),
            "source": "usuarios_xlsx",
            "raw_payload": row,
        }
        directory_rows_map[email] = payload

    legacy_items = fetch_all_legacy_items(supabase_url, service_key)

    links_map: Dict[tuple, Dict[str, Any]] = {}
    for item in legacy_items:
        legacy_year = item.get("legacy_year")
        demand_code = str(item.get("demand_code") or "").strip() or None
        email = normalize_email(item.get("fiscal_email"))
        if not legacy_year or not demand_code or not email:
            continue
        key = (legacy_year, demand_code, email, "fiscal")
        links_map[key] = {
            "legacy_year": legacy_year,
            "demand_code": demand_code,
            "user_email": email,
            "link_role": "fiscal",
            "source": "itens_solicitados_col_n",
        }
        if email not in directory_rows_map:
            directory_rows_map[email] = {
                "email": email,
                "server_name": None,
                "birth_date": None,
                "start_date": None,
                "source": "legacy_items_only",
                "raw_payload": {"origin": "legacy_pa_itens.fiscal_email"},
            }

    directory_rows = list(directory_rows_map.values())
    link_rows = list(links_map.values())

    summary = {
        "file": str(xlsx_path),
        "apply": args.apply,
        "users_sheet_rows": len(rows),
        "directory_rows_upsert": len(directory_rows),
        "directory_from_users_xlsx": sum(1 for r in directory_rows if r["source"] == "usuarios_xlsx"),
        "directory_from_legacy_items_only": sum(1 for r in directory_rows if r["source"] == "legacy_items_only"),
        "legacy_items_rows": len(legacy_items),
        "demand_user_links_upsert": len(link_rows),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))

    if not args.apply:
        return

    upsert_rows(
        supabase_url,
        service_key,
        "legacy_user_directory",
        directory_rows,
        "email",
    )
    upsert_rows(
        supabase_url,
        service_key,
        "legacy_demand_user_links",
        link_rows,
        "legacy_year,demand_code,user_email,link_role",
    )
    print("Importação de usuários legados e vínculos concluída com sucesso.")


if __name__ == "__main__":
    main()

