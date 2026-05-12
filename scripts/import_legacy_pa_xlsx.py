#!/usr/bin/env python3
"""
Importador do PERCATALICITA.xlsx para tabelas legadas no Supabase.

Modos:
- Dry-run (padrão): apenas analisa e mostra estatísticas.
- Apply: grava nas tabelas legacy_pa_demandas e legacy_pa_itens.

Uso:
  python scripts/import_legacy_pa_xlsx.py
  python scripts/import_legacy_pa_xlsx.py --apply
  python scripts/import_legacy_pa_xlsx.py --year 2025 --file "..\\LEGACY_FILES\\PERCATALICITA.xlsx" --apply
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple
from urllib.parse import urlencode

import openpyxl
import requests


def load_env_file(env_path: Path) -> Dict[str, str]:
    env: Dict[str, str] = {}
    if not env_path.exists():
        return env

    for line in env_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        env[key.strip()] = value.strip().strip('"').strip("'")
    return env


def unique_headers(raw_headers: Iterable[Any]) -> List[str]:
    seen: Dict[str, int] = {}
    result: List[str] = []
    for idx, header in enumerate(raw_headers, start=1):
        base = str(header).strip() if header is not None and str(header).strip() else f"__col_{idx}"
        if base not in seen:
            seen[base] = 1
            result.append(base)
        else:
            seen[base] += 1
            result.append(f"{base}__{seen[base]}")
    return result


def parse_timestamp(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()

    text = str(value).strip()
    if not text:
        return None

    candidates = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y",
    ]
    for fmt in candidates:
        try:
            return datetime.strptime(text, fmt).isoformat()
        except ValueError:
            pass

    return None


def parse_numeric(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    if not text:
        return None

    text = text.replace("R$", "").replace(".", "").replace(",", ".")
    match = re.search(r"-?\d+(\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def infer_year(code: str | None, fallback_year: int) -> int:
    if code:
        match = re.search(r"(20\d{2})", code)
        if match:
            return int(match.group(1))
    return fallback_year


def normalize_demand_code_from_item(pedido_codigo: str | None) -> str | None:
    if not pedido_codigo:
        return None
    code = pedido_codigo.strip()
    if not code:
        return None
    return re.sub(r"-I\d+$", "", code)


def extract_efisco_code_from_description(description: Any) -> str | None:
    if description is None:
        return None
    text = str(description).strip()
    if not text:
        return None

    match = re.match(r"^\s*\((\d{4,})\)", text)
    if match:
        return match.group(1)
    return None


def normalize_efisco_code_col_b(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    # remove espaços em volta de hífens para padronizar "48828 - 3" -> "48828-3"
    text = re.sub(r"\s*-\s*", "-", text)
    return text


def row_hash(payload: Dict[str, Any]) -> str:
    serial = json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha1(serial.encode("utf-8")).hexdigest()


def chunked(items: List[Dict[str, Any]], size: int) -> Iterable[List[Dict[str, Any]]]:
    for i in range(0, len(items), size):
        yield items[i : i + size]


def dedupe_rows_by_conflict(rows: List[Dict[str, Any]], on_conflict: str) -> List[Dict[str, Any]]:
    keys = [k.strip() for k in on_conflict.split(",") if k.strip()]
    if not keys:
        return rows

    deduped: Dict[Tuple[Any, ...], Dict[str, Any]] = {}
    passthrough: List[Dict[str, Any]] = []
    for row in rows:
        if any(k not in row for k in keys):
            passthrough.append(row)
            continue
        key = tuple(row.get(k) for k in keys)
        deduped[key] = row
    return list(deduped.values()) + passthrough


def upsert_rows(
    base_url: str,
    service_key: str,
    table: str,
    rows: List[Dict[str, Any]],
    on_conflict: str,
) -> None:
    if not rows:
        return

    endpoint = f"{base_url}/rest/v1/{table}"
    query = urlencode({"on_conflict": on_conflict})
    url = f"{endpoint}?{query}"
    headers = {
        "apikey": service_key,
        "Authorization": f"Bearer {service_key}",
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates,return=minimal",
    }

    deduped_rows = dedupe_rows_by_conflict(rows, on_conflict)
    for batch in chunked(deduped_rows, 500):
        # Requests/json cannot encode datetime objects that may appear in raw_payload.
        sanitized_batch = json.loads(json.dumps(batch, ensure_ascii=False, default=str))
        resp = requests.post(url, headers=headers, json=sanitized_batch, timeout=120)
        if resp.status_code >= 300:
            raise RuntimeError(
                f"Upsert failed on table {table} ({resp.status_code}): {resp.text}"
            )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--file",
        default=str(
            (Path(__file__).resolve().parents[2] / "LEGACY_FILES" / "PERCATALICITA.xlsx")
        ),
    )
    parser.add_argument("--year", type=int, default=2025)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    env_local = load_env_file(Path(__file__).resolve().parents[1] / ".env.local")
    supabase_url = os.getenv("NEXT_PUBLIC_SUPABASE_URL") or env_local.get("NEXT_PUBLIC_SUPABASE_URL")
    service_key = os.getenv("SUPABASE_SERVICE_ROLE_KEY") or env_local.get("SUPABASE_SERVICE_ROLE_KEY")

    if not supabase_url or not service_key:
        raise RuntimeError("NEXT_PUBLIC_SUPABASE_URL e SUPABASE_SERVICE_ROLE_KEY são obrigatórios.")

    xlsx_path = Path(args.file).resolve()
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws_demands = wb["Demandas"]
    ws_items = wb["Itens_Solicitados"]

    demand_headers = unique_headers(next(ws_demands.iter_rows(min_row=1, max_row=1, values_only=True)))
    item_headers = unique_headers(next(ws_items.iter_rows(min_row=1, max_row=1, values_only=True)))

    demands_rows: List[Dict[str, Any]] = []
    for raw in ws_demands.iter_rows(min_row=2, values_only=True):
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        data = {demand_headers[i]: raw[i] if i < len(raw) else None for i in range(len(demand_headers))}
        demand_code = str(data.get("Código Demanda") or "").strip()
        if not demand_code:
            continue

        legacy_year = infer_year(demand_code, args.year)
        payload = {
            "legacy_year": legacy_year,
            "demand_code": demand_code,
            "record_origin": "demandas_sheet",
            "submission_at": parse_timestamp(data.get("Data Submissão")),
            "request_area": data.get("Área Requisitante"),
            "requester_name": data.get("Responsável"),
            "requester_email": str(data.get("E-mail") or "").strip().lower() or None,
            "object": data.get("Objeto"),
            "current_problem": data.get("Problema atual"),
            "justification_acquisition": data.get("Justificativa para aquisição"),
            "justification_quantity": data.get("Justificativa Quantidade"),
            "proposed_solution": data.get("Solução"),
            "total_estimated": parse_numeric(data.get("Valor Estimado Geral")),
            "delivery_forecast": parse_timestamp(data.get("Data Prevista Entrega")),
            # Legado é apenas para consulta histórica: status institucional fixo.
            "status": "Em pactuação",
            "campus": data.get("Campus"),
            "priority": data.get("Prioridade_Demanda"),
            "planning_owner": data.get("Responsável_Planejamento"),
            "demand_type": data.get("Tipo de Demanda"),
            "workflow_status": data.get("Workflow Status"),
            "delivery_status": data.get("Status Entrega"),
            "raw_payload": data,
        }
        demands_rows.append(payload)

    demand_code_set = {(d["legacy_year"], d["demand_code"]) for d in demands_rows}

    items_rows: List[Dict[str, Any]] = []
    items_with_code_from_col_c = 0
    items_with_code_from_col_b = 0
    items_without_any_code = 0

    for raw in ws_items.iter_rows(min_row=2, values_only=True):
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        data = {item_headers[i]: raw[i] if i < len(raw) else None for i in range(len(item_headers))}

        pedido_codigo = str(data.get("Pedido_Código") or "").strip() or None
        demand_code = normalize_demand_code_from_item(pedido_codigo) if pedido_codigo else None
        legacy_year = infer_year(demand_code or pedido_codigo, args.year)

        efisco_code_col_b = normalize_efisco_code_col_b(data.get("Código_EFISCO"))
        efisco_code_col_c = extract_efisco_code_from_description(data.get("Descrição_EFISCO"))
        if efisco_code_col_b:
            efisco_code = efisco_code_col_b
            efisco_source = "col_b"
            items_with_code_from_col_b += 1
        elif efisco_code_col_c:
            efisco_code = efisco_code_col_c
            efisco_source = "col_c_parentheses"
            items_with_code_from_col_c += 1
        else:
            efisco_code = None
            efisco_source = "none"
            items_without_any_code += 1

        payload = {
            "legacy_year": legacy_year,
            "pedido_codigo": pedido_codigo,
            "demand_code": demand_code,
            "efisco_code": efisco_code,
            "efisco_code_col_b": efisco_code_col_b,
            "efisco_code_col_c": efisco_code_col_c,
            "efisco_code_source": efisco_source,
            "efisco_description": data.get("Descrição_EFISCO"),
            "quantity_text": str(data.get("Quantidade") or "").strip() or None,
            "quantity_numeric": parse_numeric(data.get("Quantidade")),
            "unit": data.get("Unidade"),
            "item_justification": data.get("Justificativa_Item"),
            "item_requirements": data.get("Requisitos_Item"),
            "local_uso": data.get("Local_Uso"),
            "ref_link_1": data.get("Link_Referencia1"),
            "ref_link_2": data.get("Link_Referencia2"),
            "ref_link_3": data.get("Link_Referencia3"),
            "photo_url": data.get("URL_Foto"),
            "fiscal_name": data.get("Servidor_Fiscal"),
            "fiscal_email": str(data.get("E-mail_Fiscal") or "").strip().lower() or None,
            "item_criticality": data.get("Criticidade Individual"),
            "item_delivery_deadline": str(data.get("Prazo Entrega Item") or "").strip() or None,
            "item_price": parse_numeric(data.get("Preço")),
            "curso": data.get("Curso"),
            "grupo": data.get("Grupo"),
            "natureza_despesa": data.get("Natureza de Despesa"),
            "is_reagente": str(data.get("É Reagente?") or "").strip() or None,
            "raw_payload": data,
        }
        payload["source_row_hash"] = row_hash(payload)
        items_rows.append(payload)

        if demand_code and (legacy_year, demand_code) not in demand_code_set:
            synthetic = {
                "legacy_year": legacy_year,
                "demand_code": demand_code,
                "record_origin": "item_only",
                "submission_at": None,
                "request_area": None,
                "requester_name": None,
                "requester_email": None,
                "object": "[Sem registro na aba Demandas]",
                "current_problem": None,
                "justification_acquisition": None,
                "justification_quantity": None,
                "proposed_solution": None,
                "total_estimated": None,
                "delivery_forecast": None,
                "status": "Em pactuação",
                "campus": None,
                "priority": None,
                "planning_owner": None,
                "demand_type": None,
                "workflow_status": None,
                "delivery_status": None,
                "raw_payload": {"synthetic_from_item": pedido_codigo},
            }
            demands_rows.append(synthetic)
            demand_code_set.add((legacy_year, demand_code))

    unique_fiscal_emails = sorted(
        {r["fiscal_email"] for r in items_rows if r.get("fiscal_email")}
    )
    missing_fiscal_email = sum(1 for r in items_rows if not r.get("fiscal_email"))
    missing_demand_code = sum(1 for r in items_rows if not r.get("demand_code"))

    summary = {
        "file": str(xlsx_path),
        "apply": args.apply,
        "demands_rows": len(demands_rows),
        "items_rows": len(items_rows),
        "unique_fiscal_emails": len(unique_fiscal_emails),
        "missing_fiscal_email_items": missing_fiscal_email,
        "missing_demand_code_items": missing_demand_code,
        "items_with_efisco_from_col_c": items_with_code_from_col_c,
        "items_with_efisco_from_col_b": items_with_code_from_col_b,
        "items_without_any_efisco_code": items_without_any_code,
        "sample_fiscal_emails": unique_fiscal_emails[:10],
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))

    if not args.apply:
        return

    upsert_rows(
        supabase_url,
        service_key,
        "legacy_pa_demandas",
        demands_rows,
        "legacy_year,demand_code",
    )
    upsert_rows(
        supabase_url,
        service_key,
        "legacy_pa_itens",
        items_rows,
        "source_row_hash",
    )
    print("Importação concluída com sucesso.")


if __name__ == "__main__":
    main()
